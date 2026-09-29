# 本机智能体 × Office：DSML 标记抢救 → 读 Word → 按段落修改（保留格式/图片）→ 生成 Word → 确认写入 → 撤销；设置页“联网搜索”“Tectonic”
import asyncio, sys, os, shutil, tempfile
from playwright.async_api import async_playwright
HERE=os.path.dirname(os.path.abspath(__file__))
ROOT=os.path.abspath(os.path.join(HERE,"..",".."))
B=sys.argv[1] if len(sys.argv)>1 else "http://127.0.0.1:18825"
OUT=os.path.join(HERE,"out","office")
def one(s): return s.replace("\n"," | ")
async def api(pg, method, path, body=None):
    return await pg.evaluate("""async ([m,p,b])=>{const r=await fetch(p,{method:m,headers:{'Content-Type':'application/json','X-KY':'1'},body:b?JSON.stringify(b):undefined});return [r.status, await r.json().catch(()=>null)]}""",[method,path,body])
async def main():
    os.makedirs(OUT, exist_ok=True)
    WS=tempfile.mkdtemp(prefix="cando-ws-")
    shutil.copy(os.path.join(ROOT,"testdata","office","report.docx"), WS)
    orig=open(os.path.join(WS,"report.docx"),"rb").read()
    async with async_playwright() as p:
        b=await p.chromium.launch(); errs=[]
        ctx=await b.new_context(viewport={"width":1360,"height":900})
        pg=await ctx.new_page(); pg.on("pageerror", lambda e: errs.append(str(e)))
        pg.on("dialog", lambda d: asyncio.ensure_future(d.accept()))
        await pg.goto(B)
        await api(pg,"POST","/api/setup",{"org_name":"智能感知课题组","username":"admin","name":"陈老师","password":"teach123"})
        await api(pg,"POST","/api/login",{"username":"admin","password":"teach123"})
        st,m=await api(pg,"POST","/api/models",{"name":"DeepSeek","base_url":"http://127.0.0.1:11434/v1","model":"deepseek-v4-pro","key":"k"})
        mid=m["mine"][0]["id"]; await api(pg,"POST",f"/api/models/{mid}/check"); await api(pg,"POST",f"/api/models/{mid}/activate",{"active":True})
        await api(pg,"PUT","/api/agent/folders",{"folders":[{"path":WS,"write":True}]})
        await pg.goto(B+"/#/assistant/agent"); await pg.reload(); await pg.wait_for_selector("#agText")
        await pg.fill("#agText","帮我改一下 Word：摘要里补上研究方法，引言后加一段研究缺口，再写一份修改说明"); await pg.click("#agGo")
        await pg.wait_for_selector("#agSideBody .change", timeout=30000)
        await pg.wait_for_function("!document.querySelector('#agGo').disabled", timeout=30000)
        print("STEPS:", one(await pg.inner_text("#agSteps"))[:500] if await pg.query_selector("#agSteps") else "")
        chs=await pg.locator("#agSideBody .change").all_inner_texts()
        print("CHANGES:", [one(c)[:120] for c in chs])
        await pg.screenshot(path=f"{OUT}/agent_word.png")
        await pg.click("button[data-act=agApplyAll]"); await pg.wait_for_timeout(1500)
        import docx
        d=docx.Document(os.path.join(WS,"report.docx"))
        print("AFTER APPLY:", [p.text for p in d.paragraphs][:6], "| images:", sum(1 for r in d.part.rels.values() if "image" in r.reltype), "| tables:", len(d.tables))
        print("NEW DOC:", os.path.exists(os.path.join(WS,"修改说明.docx")))
        m2=docx.Document(os.path.join(WS,"修改说明.docx")); print("NEW DOC TEXT:", [p.text for p in m2.paragraphs], "tables:", len(m2.tables))
        await pg.screenshot(path=f"{OUT}/agent_applied.png")
        # 撤销 Word 修改
        for i in range(await pg.locator("#agSideBody .change").count()):
            c=pg.locator("#agSideBody .change").nth(i)
            if "修改 Word" in await c.inner_text():
                await c.locator("button[data-a=undo]").click(); break
        await pg.wait_for_timeout(800)
        print("UNDO restores original:", open(os.path.join(WS,"report.docx"),"rb").read()==orig)
        # 设置：LaTeX 与联网搜索
        await pg.click("button[data-act=agSide][data-s=folders]"); await pg.wait_for_selector("form[data-submit=agPolicy]")
        print("TEX SECTION:", one(await pg.inner_text("#agSideBody"))[-260:])
        await pg.goto(B+"/#/settings"); await pg.wait_for_selector("#webCard form[data-submit=webKeySave]")
        print("WEB CARD:", one(await pg.inner_text("#webCard"))[:160])
        await pg.locator("#webCard").screenshot(path=f"{OUT}/websearch_card.png")
        print("ERRORS:", errs); await b.close()
asyncio.run(main())
